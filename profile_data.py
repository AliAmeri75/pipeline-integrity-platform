"""Read grid-based crack and corrosion profiles for the Data Integration page."""

from __future__ import annotations

import io
import math
import random
import warnings
from datetime import time
from statistics import median
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


SUMMARY_COLUMNS = [
    "Profile ID",
    "GW ID",
    "Feature ID",
    "Defect type",
    "Feature location (m)",
    "Profile start (m)",
    "Profile end (m)",
    "Clock position",
    "Maximum depth (%WT)",
    "Maximum depth (mm)",
    "Profile length (mm)",
    "Profile width (mm)",
    "Wall thickness (mm)",
    "Outside diameter (mm)",
    "Measured profile cells",
]


class ProfileFormatError(ValueError):
    """Raised when a workbook does not contain recognizable grid profiles."""


def synthetic_sample_workbook(seed: int = 7314) -> bytes:
    """Create a reproducible workbook containing sample crack profiles."""

    random_generator = random.Random(seed)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "SAMPLE DATA FOR CRACK PROFILE"
    summary["A1"].font = Font(bold=True, color="9C0006")
    summary["A1"].fill = PatternFill("solid", fgColor="FFC7CE")
    summary.merge_cells("A1:I1")
    summary.append([])
    summary_headers = [
        "File Name",
        "GW ID",
        "CLS ID",
        "TCPL\nChainage\n(m)",
        "Clock\nPosition\n(hh:mm)",
        "Outside Diameter (mm)",
        "Wall Thickness (mm)",
        "Peak Depth (%wt)",
        "Length (mm)",
    ]
    summary.append(summary_headers)
    summary.freeze_panes = "A4"
    summary.sheet_view.showGridLines = False

    specifications = [
        {
            "profile_id": "SYN-001",
            "gw_id": "DEMO-GW-01",
            "feature_id": "DEMO-F-101",
            "chainage_m": 500.0,
            "clock": "02:15",
            "diameter_mm": 609.6,
            "wall_mm": 9.5,
            "x_count": 36,
            "y_count": 17,
            "x_step": 2.0,
            "y_step": 4.0,
            "peak_target": 32.0,
        },
        {
            "profile_id": "SYN-002",
            "gw_id": "DEMO-GW-02",
            "feature_id": "DEMO-F-202",
            "chainage_m": 4500.0,
            "clock": "06:40",
            "diameter_mm": 762.0,
            "wall_mm": 10.3,
            "x_count": 48,
            "y_count": 21,
            "x_step": 2.0,
            "y_step": 4.0,
            "peak_target": 44.0,
        },
        {
            "profile_id": "SYN-003",
            "gw_id": "DEMO-GW-03",
            "feature_id": "DEMO-F-303",
            "chainage_m": 8500.0,
            "clock": "10:05",
            "diameter_mm": 914.4,
            "wall_mm": 11.8,
            "x_count": 60,
            "y_count": 19,
            "x_step": 6.0,
            "y_step": 4.0,
            "peak_target": 53.0,
        },
    ]

    for index, specification in enumerate(specifications):
        profile_id = str(specification["profile_id"])
        feature_id = str(specification["feature_id"])
        chainage_m = float(specification["chainage_m"])
        x_count = int(specification["x_count"])
        y_count = int(specification["y_count"])
        x_step = float(specification["x_step"])
        y_step = float(specification["y_step"])
        peak_target = float(specification["peak_target"])
        profile_length = x_count * x_step
        start_mm = chainage_m * 1000.0
        circumferential_start = 200.0 + 140.0 * index

        worksheet = workbook.create_sheet(profile_id)
        worksheet.sheet_view.showGridLines = False
        worksheet["A1"] = "SAMPLE DATA FOR CRACK PROFILE"
        worksheet["A1"].font = Font(bold=True, color="9C0006")
        worksheet["E1"] = (
            "Columns: Longitudinal Coordinate (mm) / "
            "Rows: Circumferential Coordinate (mm) / Cells: Crack Depth (%WT)"
        )
        worksheet["E1"].font = Font(bold=True, color="FFFFFF")
        worksheet["E1"].fill = PatternFill("solid", fgColor="275D38")

        axial_coordinates = [start_mm + x_step * offset for offset in range(x_count)]
        circumferential_coordinates = [
            circumferential_start + y_step * offset for offset in range(y_count)
        ]
        for column_number, coordinate in enumerate(axial_coordinates, start=6):
            worksheet.cell(1, column_number, coordinate)
        for row_number, coordinate in enumerate(circumferential_coordinates, start=2):
            worksheet.cell(row_number, 5, coordinate)

        centers = [
            (
                random_generator.uniform(0.30, 0.48) * (x_count - 1),
                random_generator.uniform(0.35, 0.60) * (y_count - 1),
                peak_target,
                random_generator.uniform(5.0, 8.0),
                random_generator.uniform(2.5, 4.5),
            ),
            (
                random_generator.uniform(0.58, 0.78) * (x_count - 1),
                random_generator.uniform(0.28, 0.72) * (y_count - 1),
                peak_target * random_generator.uniform(0.48, 0.70),
                random_generator.uniform(4.0, 7.0),
                random_generator.uniform(2.0, 4.0),
            ),
        ]
        maximum_depth = 0.0
        for y_index, _ in enumerate(circumferential_coordinates):
            for x_index, _ in enumerate(axial_coordinates):
                depth = 0.0
                for center_x, center_y, amplitude, spread_x, spread_y in centers:
                    exponent = -0.5 * (
                        ((x_index - center_x) / spread_x) ** 2
                        + ((y_index - center_y) / spread_y) ** 2
                    )
                    depth += amplitude * math.exp(exponent)
                depth += random_generator.uniform(-1.25, 1.25)
                depth = 0.0 if depth < 4.0 else round(min(depth, 65.0), 1)
                maximum_depth = max(maximum_depth, depth)
                worksheet.cell(y_index + 2, x_index + 6, depth)

        worksheet["B4"] = "Feature ID"
        worksheet["C4"] = feature_id
        worksheet["B5"] = "Outside Diameter (mm)"
        worksheet["C5"] = float(specification["diameter_mm"])
        worksheet["B6"] = "Wall Thickness (mm)"
        worksheet["C6"] = float(specification["wall_mm"])
        worksheet["B18"] = "Peak Depth (%wt)"
        worksheet["C18"] = maximum_depth
        worksheet["B19"] = "Total Length (mm)"
        worksheet["C19"] = profile_length
        worksheet["A24"] = "Sample data for crack profile."
        worksheet["A24"].font = Font(italic=True, color="9C0006")
        worksheet.column_dimensions["A"].width = 34
        worksheet.column_dimensions["B"].width = 25
        worksheet.column_dimensions["C"].width = 18
        worksheet.column_dimensions["E"].width = 18

        summary.append(
            [
                profile_id,
                specification["gw_id"],
                feature_id,
                chainage_m,
                specification["clock"],
                specification["diameter_mm"],
                specification["wall_mm"],
                maximum_depth,
                profile_length,
            ]
        )

    for column_number, header in enumerate(summary_headers, start=1):
        cell = summary.cell(3, column_number, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="275D38")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        summary.column_dimensions[cell.column_letter].width = 20
    summary.auto_filter.ref = f"A3:I{3 + len(specifications)}"

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _clean_text(value: Any) -> str:
    return " ".join(str(value).replace("\n", " ").replace("_", " ").split()).strip()


def _normalized(value: Any) -> str:
    return _clean_text(value).casefold()


def _number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return None
    return numeric if math.isfinite(numeric) else None


def _identifier(value: Any) -> str:
    numeric = _number(value)
    if numeric is not None and numeric.is_integer():
        return str(int(numeric))
    return _clean_text(value)


def _clock_text(value: Any) -> str:
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return _clean_text(value)


def _find_label_value(worksheet: Any, label: str) -> Any:
    """Return the value immediately to the right of a matching label."""

    target = _normalized(label)
    for row in worksheet.iter_rows(
        min_row=1,
        max_row=min(worksheet.max_row, 40),
        min_col=1,
        max_col=min(worksheet.max_column, 6),
    ):
        for cell in row:
            if _normalized(cell.value) == target:
                return worksheet.cell(cell.row, cell.column + 1).value
    return None


def _summary_records(workbook: Any) -> dict[str, dict[str, Any]]:
    """Read the optional TCAT summary and index it by profile file name."""

    summary_name = next(
        (name for name in workbook.sheetnames if name.casefold() == "summary"), None
    )
    if summary_name is None:
        return {}

    worksheet = workbook[summary_name]
    header_row = None
    headers: list[str] = []
    for row_number in range(1, min(worksheet.max_row, 20) + 1):
        values = [worksheet.cell(row_number, column).value for column in range(1, worksheet.max_column + 1)]
        normalized = [_normalized(value) for value in values]
        if "file name" in normalized and any("feature" in value or "cls id" in value for value in normalized):
            header_row = row_number
            headers = normalized
            break
    if header_row is None:
        return {}

    records: dict[str, dict[str, Any]] = {}
    for row_number in range(header_row + 1, worksheet.max_row + 1):
        values = [worksheet.cell(row_number, column).value for column in range(1, len(headers) + 1)]
        record = {header: value for header, value in zip(headers, values) if header}
        file_name = record.get("file name")
        if file_name not in (None, ""):
            records[_identifier(file_name)] = record
    return records


def _summary_value(record: dict[str, Any], *candidates: str) -> Any:
    for candidate in candidates:
        key = _normalized(candidate)
        if key in record and record[key] not in (None, ""):
            return record[key]
    return None


def _positive_spacing(values: list[float]) -> float:
    unique = sorted(set(values))
    differences = [right - left for left, right in zip(unique, unique[1:]) if right > left]
    return float(median(differences)) if differences else 0.0


def _profile_header(worksheet: Any) -> tuple[int, int] | None:
    for row in worksheet.iter_rows(
        min_row=1,
        max_row=min(worksheet.max_row, 20),
        min_col=1,
        max_col=min(worksheet.max_column, 12),
    ):
        for cell in row:
            text = _normalized(cell.value)
            if "longitudinal coordinate" in text and "depth" in text:
                return cell.row, cell.column
    return None


def _contiguous_axis(worksheet: Any, *, row: int | None = None, column: int | None = None, start: int) -> list[tuple[int, float]]:
    axis: list[tuple[int, float]] = []
    limit = worksheet.max_column if row is not None else worksheet.max_row
    for index in range(start, limit + 1):
        value = worksheet.cell(row, index).value if row is not None else worksheet.cell(index, column).value
        numeric = _number(value)
        if numeric is None:
            if axis:
                break
            continue
        axis.append((index, numeric))
    return axis


def _extract_profile(worksheet: Any, summary_record: dict[str, Any]) -> tuple[dict[str, Any], pd.DataFrame] | None:
    header = _profile_header(worksheet)
    if header is None:
        return None

    header_row, header_column = header
    axial_axis = _contiguous_axis(
        worksheet, row=header_row, start=header_column + 1
    )
    circumferential_axis = _contiguous_axis(
        worksheet, column=header_column, start=header_row + 1
    )
    if not axial_axis or not circumferential_axis:
        return None

    axial_values = [value for _, value in axial_axis]
    circumferential_values = [value for _, value in circumferential_axis]
    axial_spacing = _positive_spacing(axial_values)
    circumferential_spacing = _positive_spacing(circumferential_values)
    profile_id = _identifier(worksheet.title)
    feature_id = _identifier(
        _find_label_value(worksheet, "Feature ID")
        or _summary_value(summary_record, "CLS ID", "Feature ID")
        or profile_id
    )

    records: list[dict[str, Any]] = []
    measured_depths: list[float] = []
    for row_number, circumferential_position in circumferential_axis:
        for column_number, axial_position in axial_axis:
            depth = _number(worksheet.cell(row_number, column_number).value)
            if depth is None:
                continue
            measured_depths.append(depth)
            axial_offset = axial_position - min(axial_values)
            circumferential_offset = circumferential_position - min(circumferential_values)
            records.append(
                {
                    "Profile ID": profile_id,
                    "Feature ID": feature_id,
                    "Axial location (mm)": axial_position,
                    "Axial offset (mm)": axial_offset,
                    "Axial cell start (mm)": axial_offset - axial_spacing / 2.0,
                    "Axial cell end (mm)": axial_offset + axial_spacing / 2.0,
                    "Circumferential location (mm)": circumferential_position,
                    "Circumferential offset (mm)": circumferential_offset,
                    "Circumferential cell start (mm)": circumferential_offset - circumferential_spacing / 2.0,
                    "Circumferential cell end (mm)": circumferential_offset + circumferential_spacing / 2.0,
                    "Depth (%WT)": depth,
                }
            )
    if not measured_depths:
        return None

    maximum_depth_pct = max(measured_depths)
    reported_length = _number(_find_label_value(worksheet, "Total Length (mm)"))
    reported_width = _number(_find_label_value(worksheet, "Total Width (mm)"))
    derived_length = (
        max(axial_values) - min(axial_values) + axial_spacing
        if len(axial_values) > 1
        else 0.0
    )
    derived_width = (
        max(circumferential_values) - min(circumferential_values) + circumferential_spacing
        if len(circumferential_values) > 1
        else 0.0
    )
    profile_length = reported_length if reported_length and reported_length > 0 else derived_length
    profile_width = reported_width if reported_width and reported_width > 0 else derived_width

    wall_thickness = _number(_find_label_value(worksheet, "Wall Thickness (mm)"))
    outside_diameter = _number(_find_label_value(worksheet, "Outside Diameter (mm)"))
    summary_location = _number(
        _summary_value(summary_record, "TCPL Chainage (m)", "Feature Chainage (m)")
    )
    profile_start = min(axial_values) / 1000.0
    feature_location = summary_location if summary_location is not None else profile_start
    profile_end = profile_start + profile_length / 1000.0

    metadata = {
        "Profile ID": profile_id,
        "GW ID": _identifier(_summary_value(summary_record, "GW ID")),
        "Feature ID": feature_id,
        "Defect type": "Corrosion / metal loss",
        "Feature location (m)": feature_location,
        "Profile start (m)": profile_start,
        "Profile end (m)": profile_end,
        "Clock position": _clock_text(
            _summary_value(summary_record, "Clock Position (hh:mm)")
        ),
        "Maximum depth (%WT)": maximum_depth_pct,
        "Maximum depth (mm)": (
            wall_thickness * maximum_depth_pct / 100.0
            if wall_thickness is not None
            else None
        ),
        "Profile length (mm)": profile_length,
        "Profile width (mm)": profile_width,
        "Wall thickness (mm)": wall_thickness,
        "Outside diameter (mm)": outside_diameter,
        "Measured profile cells": len(measured_depths),
    }
    return metadata, pd.DataFrame.from_records(records)


def parse_profile_workbook(content: bytes) -> tuple[pd.DataFrame, dict[str, pd.DataFrame]]:
    """Parse all recognizable TCAT/Psqr grid-profile sheets in an Excel workbook."""

    try:
        # Normal worksheet mode is substantially faster for these compact but wide
        # profile grids. The in-memory workbook is never saved back to the upload.
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="DrawingML support is incomplete.*")
            warnings.filterwarnings("ignore", message="Data Validation extension is not supported.*")
            workbook = load_workbook(
                io.BytesIO(content), read_only=False, data_only=True, keep_vba=False
            )
    except Exception as error:  # openpyxl raises several format-specific errors
        raise ProfileFormatError(f"The uploaded file could not be read as an Excel workbook: {error}") from error

    summaries: list[dict[str, Any]] = []
    profiles: dict[str, pd.DataFrame] = {}
    try:
        summary_records = _summary_records(workbook)
        for sheet_name in workbook.sheetnames:
            sheet_key = _identifier(sheet_name)
            if summary_records and sheet_key not in summary_records:
                continue
            if not summary_records and any(
                word in sheet_key.casefold()
                for word in ("master", "example", "template")
            ):
                continue
            record = summary_records.get(sheet_key, {})
            extracted = _extract_profile(workbook[sheet_name], record)
            if extracted is None:
                continue
            metadata, profile = extracted
            profile_id = str(metadata["Profile ID"])
            summaries.append(metadata)
            profiles[profile_id] = profile
    finally:
        workbook.close()

    if not summaries:
        raise ProfileFormatError(
            "No grid profiles were found. A profile sheet needs longitudinal coordinates "
            "across the top, circumferential coordinates down the side, and depth values "
            "inside the grid."
        )

    summary = pd.DataFrame.from_records(summaries, columns=SUMMARY_COLUMNS)
    return summary.sort_values("Feature location (m)", kind="stable").reset_index(drop=True), profiles


def maximum_depth_envelope(profile: pd.DataFrame) -> pd.DataFrame:
    """Return maximum measured depth at each axial profile position."""

    return (
        profile.groupby("Axial offset (mm)", as_index=False)["Depth (%WT)"]
        .max()
        .rename(columns={"Depth (%WT)": "Maximum depth (%WT)"})
        .sort_values("Axial offset (mm)")
    )


def summary_for_defect_type(summary: pd.DataFrame, defect_type: str) -> pd.DataFrame:
    """Return the planning fields appropriate for a crack or corrosion profile."""

    output = summary.copy()
    output["Defect type"] = defect_type
    if defect_type == "Crack":
        output = output.drop(columns=["Profile width (mm)"], errors="ignore")
    return output


def summary_csv(summary: pd.DataFrame) -> bytes:
    return summary.to_csv(index=False, float_format="%.6g").encode("utf-8-sig")


def summary_xlsx(summary: pd.DataFrame) -> bytes:
    """Create a small, readable Excel handoff for inspection-planning inputs."""

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Profile Summary"
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False

    for column_number, column_name in enumerate(summary.columns, start=1):
        cell = worksheet.cell(1, column_number, column_name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="275D38")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_number, values in enumerate(summary.itertuples(index=False, name=None), start=2):
        for column_number, value in enumerate(values, start=1):
            worksheet.cell(row_number, column_number, None if pd.isna(value) else value)

    numeric_formats = {
        "Feature location (m)": "0.000",
        "Profile start (m)": "0.000",
        "Profile end (m)": "0.000",
        "Maximum depth (%WT)": "0.00",
        "Maximum depth (mm)": "0.000",
        "Profile length (mm)": "0.00",
        "Profile width (mm)": "0.00",
        "Wall thickness (mm)": "0.000",
        "Outside diameter (mm)": "0.0",
        "Measured profile cells": "0",
    }
    for column_number, column_name in enumerate(summary.columns, start=1):
        width = max(
            len(str(column_name)),
            *(len(str(value)) for value in summary[column_name].head(200) if not pd.isna(value)),
        )
        worksheet.column_dimensions[worksheet.cell(1, column_number).column_letter].width = min(max(width + 2, 12), 25)
        if column_name in numeric_formats:
            for row_number in range(2, len(summary) + 2):
                worksheet.cell(row_number, column_number).number_format = numeric_formats[column_name]

    worksheet.auto_filter.ref = worksheet.dimensions
    if len(summary) > 0:
        table = Table(displayName="ProfileSummary", ref=worksheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
