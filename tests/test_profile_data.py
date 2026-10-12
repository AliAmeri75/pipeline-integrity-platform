import io
import unittest

from openpyxl import Workbook, load_workbook

from profile_data import (
    maximum_depth_envelope,
    parse_profile_workbook,
    summary_csv,
    summary_xlsx,
    synthetic_sample_workbook,
)


def _sample_workbook() -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Vendor Tally"])
    summary.append(
        [
            "File Name",
            "GW ID",
            "CLS ID",
            "TCPL\nChainage\n(m)",
            "Clock\nPosition\n(hh:mm)",
        ]
    )
    summary.append(["TEST-001", "TEST-GW-01", "TEST-F-01", 80.0, "07:51"])

    profile = workbook.create_sheet("TEST-001")
    profile["E1"] = "Longitudinal Coordinate (mm) -> / Circumferential Coordinate (mm) / Depth (%WT)"
    for column, value in enumerate([80000, 80001, 80002], start=6):
        profile.cell(1, column, value)
    for row, circumferential in enumerate([100.0, 105.0], start=2):
        profile.cell(row, 5, circumferential)
    for row, values in enumerate([[0, 20, 40], [10, 50, 30]], start=2):
        for column, value in enumerate(values, start=6):
            profile.cell(row, column, value)
    profile["B4"] = "Feature ID"
    profile["C4"] = "TEST-F-01"
    profile["B5"] = "Outside Diameter (mm)"
    profile["C5"] = 863.6
    profile["B6"] = "Wall Thickness (mm)"
    profile["C6"] = 8.0
    profile["B18"] = "Peak Depth (%wt)"
    profile["C18"] = 50
    profile["B19"] = "Total Length (mm)"
    profile["C19"] = 3
    profile["B20"] = "Total Width (mm)"
    profile["C20"] = 10

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class ProfileDataTests(unittest.TestCase):
    def test_built_in_sample_is_synthetic_and_parseable(self):
        sample = synthetic_sample_workbook()
        workbook = load_workbook(io.BytesIO(sample), data_only=True)
        self.assertIn("SYNTHETIC", workbook["Summary"]["A1"].value)
        self.assertTrue(all("SYN" in name for name in workbook.sheetnames if name != "Summary"))
        workbook.close()

        summary, profiles = parse_profile_workbook(sample)
        self.assertEqual(len(summary), 3)
        self.assertEqual(set(summary["Profile ID"]), {"SYN-001", "SYN-002", "SYN-003"})
        self.assertEqual(len(profiles), 3)
        self.assertTrue(summary["GW ID"].str.startswith("DEMO-GW-").all())
        self.assertTrue(summary["Maximum depth (%WT)"].between(20.0, 65.0).all())
        self.assertTrue((summary["Profile length (mm)"] > 0).all())

    def test_psqr_grid_becomes_one_planning_row(self):
        summary, profiles = parse_profile_workbook(_sample_workbook())
        self.assertEqual(len(summary), 1)
        self.assertEqual(summary.loc[0, "Profile ID"], "TEST-001")
        self.assertEqual(summary.loc[0, "Feature ID"], "TEST-F-01")
        self.assertEqual(summary.loc[0, "GW ID"], "TEST-GW-01")
        self.assertAlmostEqual(summary.loc[0, "Maximum depth (%WT)"], 50.0)
        self.assertAlmostEqual(summary.loc[0, "Maximum depth (mm)"], 4.0)
        self.assertAlmostEqual(summary.loc[0, "Profile length (mm)"], 3.0)
        self.assertAlmostEqual(summary.loc[0, "Profile width (mm)"], 10.0)
        self.assertEqual(len(profiles["TEST-001"]), 6)

    def test_axial_envelope_keeps_maximum_depth_per_position(self):
        _, profiles = parse_profile_workbook(_sample_workbook())
        envelope = maximum_depth_envelope(profiles["TEST-001"])
        self.assertEqual(envelope["Maximum depth (%WT)"].tolist(), [10.0, 50.0, 40.0])

    def test_csv_and_excel_exports_are_readable(self):
        summary, _ = parse_profile_workbook(_sample_workbook())
        csv_text = summary_csv(summary).decode("utf-8-sig")
        self.assertIn("Maximum depth (%WT)", csv_text)
        workbook = load_workbook(io.BytesIO(summary_xlsx(summary)), data_only=True)
        worksheet = workbook["Profile Summary"]
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(worksheet["A2"].value, "TEST-001")
        self.assertEqual(len(worksheet.tables), 1)


if __name__ == "__main__":
    unittest.main()
